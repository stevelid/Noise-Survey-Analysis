"""Window-only, sequential source readers for offline annotation exports.

The dashboard's lazy loader and parsed-file cache intentionally are not used:
they retain complete meter logs. These readers scan one source in bounded chunks
and yield only samples intersecting requested source-time windows. Header and
column normalization reuse the instrument parsers without calling eager parse().
"""
from __future__ import annotations

from dataclasses import dataclass
import csv
from itertools import islice
import math
from pathlib import Path
import re
from typing import Iterator

import numpy as np
import pandas as pd

from .annotation_metrics import datetime_milliseconds, union_intervals
from .data_parsers import NoiseParserFactory, SvanFileParser, NTiFileParser, NoiseSentryFileParser

DEFAULT_CHUNK_ROWS = 16384
HEADER_SCAN_LINES = 256


@dataclass
class WindowChunk:
    frame: pd.DataFrame
    resolution: str
    sample_interval_s: float | None
    source_path: str


def _canonical_column(name: str, parser) -> str:
    name = str(name).strip().replace("_dt", "").replace(".0%", "")
    if isinstance(parser, SvanFileParser):
        mapped, _ = parser._map_svan_column(name)
        if mapped:
            return mapped
        simple = {"LAeq (TH) [dB]": "LAeq", "LAFmax (TH) [dB]": "LAFmax",
                  "LAeq Histogram (SR) [dB] L10": "LAF10", "LAeq Histogram (SR) [dB] L90": "LAF90"}
        if name in simple:
            return simple[name]
    aliases = {"start time": "Datetime", "date & time": "Datetime", "start date & time": "Datetime",
               "timestamp": "Datetime", "datetime": "Datetime", "time (date hh:mm:ss.ms)": "Datetime",
               "leq db-a": "LAeq", "leq db -a": "LAeq", "lmax db-a": "LAFmax",
               "l-max db-a": "LAFmax", "l-max db -a": "LAFmax", "l10 db-a": "LAF10",
               "l90 db-a": "LAF90", "la90": "LAF90", "la10": "LAF10"}
    canonical = aliases.get(name.lower())
    if canonical:
        return canonical
    for standard in ("LAeq", "LAFmax", "LAF90", "LAF10", "LAF", "LAFp"):
        if name.lower() == standard.lower():
            return standard
    return name


def _normalize(frame: pd.DataFrame, parser, options: dict) -> pd.DataFrame:
    frame = frame.rename(columns={column: _canonical_column(str(column), parser) for column in frame.columns})
    frame = frame.loc[:, ~frame.columns.duplicated()].copy()
    if "Datetime" not in frame:
        pair = next((names for names in (("Date", "Time"), ("Start Date", "Start Time")) if all(name in frame for name in names)), None)
        if pair:
            frame["Datetime"] = frame[pair[0]].astype(str) + " " + frame[pair[1]].astype(str)
        else:
            raise ValueError("No supported datetime column in source")
    raw = frame["Datetime"]
    if pd.api.types.is_numeric_dtype(raw):
        frame["Datetime"] = pd.to_datetime(raw, unit="ms", errors="coerce", utc=True)
    else:
        # Preserve explicit offsets/UTC. Existing parser localization is reused
        # for naive instrument wall-clock times, retaining its DST safeguards.
        text = raw.astype(str)
        explicit_zone = text.str.contains(r"(?:Z|[+-]\d{2}:?\d{2})$", regex=True)
        if explicit_zone.all():
            frame["Datetime"] = pd.to_datetime(text, errors="coerce", utc=True, format="mixed")
        elif explicit_zone.any():
            parsed = pd.Series(pd.NaT, index=frame.index, dtype="datetime64[ns, UTC]")
            parsed.loc[explicit_zone] = pd.to_datetime(text.loc[explicit_zone], errors="coerce", utc=True, format="mixed")
            naive = frame.loc[~explicit_zone, ["Datetime"]].copy()
            naive["__source_row"] = naive.index
            naive = parser._normalize_datetime_column(naive, ["Datetime"])
            if not naive.empty:
                parsed.loc[naive["__source_row"].to_list()] = naive["Datetime"].to_numpy()
            frame["Datetime"] = parsed
        else:
            frame = parser._normalize_datetime_column(frame, ["Datetime"])
    frame = frame.dropna(subset=["Datetime"])
    selected = options.get("selected_columns") or options.get("columns")
    if isinstance(selected, str):
        selected = [selected]
    if selected:
        canonical = {_canonical_column(name, parser) for name in selected}
        frame = frame[[name for name in frame if name == "Datetime" or name in canonical]]
    # Numeric conversion deliberately drops helper Date/Time strings.
    for column in list(frame.columns):
        if column != "Datetime":
            frame[column] = pd.to_numeric(frame[column], errors="coerce")
    return frame.sort_values("Datetime").reset_index(drop=True)


def _text_layout(path: Path, parser) -> tuple[list[str], int, str, str]:
    with path.open("r", encoding="utf-8-sig", errors="replace") as handle:
        head = list(islice(handle, HEADER_SCAN_LINES))
    if not head:
        raise ValueError("Empty source")
    profile = "log" if "_log" in path.name.lower() else "overview" if any(token in path.name.lower() for token in ("_summary", "_report", "_rpt")) else "unknown"
    if isinstance(parser, NTiFileParser):
        marker = next((i for i, line in enumerate(head) if line.lstrip().startswith(("# RTA Results", "# RTA LOG Results", "# Broadband Results", "# Broadband LOG Results"))), None)
        if marker is None:
            raise ValueError("NTi table header not found within bounded header scan")
        table = [(i, line) for i, line in enumerate(head) if i > marker and line.strip()]
        spectral = "RTA" in head[marker]
        names, data_index = parser._get_nti_headers_and_data_start([line for _, line in table], spectral, head, marker)
        if not names or data_index >= len(table):
            raise ValueError("NTi table header could not be normalized")
        return names, table[data_index][0], "\t", "log" if "LOG" in head[marker] else "overview"
    if isinstance(parser, SvanFileParser):
        index = next((i for i, line in enumerate(head) if "date & time" in line.lower()), None)
        if index is not None and index >= 2:
            rows = [next(csv.reader([line])) for line in head[index - 2:index + 1]]
            width = max(map(len, rows))
            names = ["_".join(row[i].strip() for row in rows if i < len(row) and row[i].strip()) or f"Unnamed_{i}" for i in range(width)]
            return parser._make_unique_headers(names), index + 1, ",", parser._get_data_profile_heuristic(head, str(path))
    for index, line in enumerate(head):
        lower = line.lower()
        if any(token in lower for token in ("datetime", "timestamp", "start time", "date & time", "time (date", "date\t", "date,")):
            delimiter = "\t" if line.count("\t") > line.count(",") else ";" if line.count(";") > line.count(",") else ","
            names = next(csv.reader([line], delimiter=delimiter))
            return names, index + 1, delimiter, profile
    raise ValueError("Supported table header not found within bounded header scan")


def _raw_frames(path: Path, parser, chunk_rows: int) -> tuple[Iterator[pd.DataFrame], str]:
    if path.suffix.lower() == ".xls":
        raise ValueError("Legacy .xls is not supported by the bounded exporter; save as .xlsx or text")
    if path.suffix.lower() == ".xlsx":
        from openpyxl import load_workbook

        def spreadsheet_frames():
            workbook = load_workbook(path, read_only=True, data_only=True)
            try:
                rows = iter(workbook.active.iter_rows(values_only=True))
                # Instrument XLSX exports can have a short title/preamble above
                # their header. Inspect a bounded prefix, then chain unused rows.
                prefix = list(islice(rows, HEADER_SCAN_LINES))
                header_index = next((index for index, row in enumerate(prefix)
                    if any(value is not None and str(value).strip().lower() in
                           ("datetime", "timestamp", "date & time", "start time", "date", "start date") for value in row)), None)
                if header_index is None:
                    raise ValueError("Spreadsheet datetime header not found within bounded scan")
                headers = [str(value).replace("\n", " ").strip() if value is not None else f"Unnamed_{index}"
                           for index, value in enumerate(prefix[header_index])]
                from itertools import chain
                rows = chain(prefix[header_index + 1:], rows)
                while batch := list(islice(rows, chunk_rows)):
                    yield pd.DataFrame(batch, columns=headers)
            finally:
                workbook.close()
        return spreadsheet_frames(), "unknown"
    names, skiprows, separator, profile = _text_layout(path, parser)
    frames = pd.read_csv(path, names=names, header=None, skiprows=skiprows, sep=separator,
                         chunksize=chunk_rows, comment="#", encoding="utf-8-sig", encoding_errors="replace",
                         on_bad_lines="warn", low_memory=False)
    return iter(frames), profile


def iter_source_windows(path: str | Path, windows, *, options: dict | None = None,
                        chunk_rows: int = DEFAULT_CHUNK_ROWS) -> Iterator[WindowChunk]:
    """Scan a source once without caching it; yield only overlapping rows.

    This is intentionally an iterator so callers release each source/chunk before
    opening the next meter. Even non-intersecting chunks are bounded in memory.
    """
    options = options or {}
    path = Path(path)
    requested = union_intervals(windows)
    if not requested:
        return
    parser = NoiseParserFactory.get_parser(str(path), options.get("parser_type", "auto"), options.get("timezone"))
    if parser is None or parser.__class__.__name__ == "AudioFileParser":
        raise ValueError(f"Unsupported acoustic source: {path}")
    frames, inferred_profile = _raw_frames(path, parser, chunk_rows)
    configured_nominal = options.get("sample_period_seconds") or options.get("sample_interval_s")
    nominal = float(configured_nominal) if configured_nominal is not None else None
    for raw in frames:
        frame = _normalize(raw, parser, options)
        if frame.empty:
            continue
        estimate = parser._calculate_sample_period(frame)
        nominal = nominal or estimate
        if nominal is None or not math.isfinite(nominal) or nominal <= 0:
            raise ValueError("Sample interval unavailable; specify sample_interval_s in source options")
        profile = options.get("data_profile") or options.get("profile") or inferred_profile
        if profile not in ("log", "overview"):
            profile = "overview" if nominal > 60 else "log"
        times = datetime_milliseconds(frame["Datetime"])
        keep = np.zeros(len(frame), dtype=bool)
        for start, end in requested:
            keep |= (times < end) & (times + nominal * 1000 > start)
        if keep.any():
            selected = frame.loc[keep].copy()
            # Preserve per-file period when one meter combines 100 ms and 1 s
            # logs; selecting only windows must not turn intervening gaps into
            # a falsely inferred sample period later in the metric engine.
            selected["__sample_interval_s"] = float(nominal)
            for column in frame.columns:
                if column != "Datetime" and not column.startswith("__"):
                    selected[f"__sample_interval_s:{column}"] = float(nominal)
            selected.attrs["sample_interval_s"] = float(nominal)
            yield WindowChunk(selected, profile, float(nominal), str(path))
